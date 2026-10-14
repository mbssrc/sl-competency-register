"""Excel import module: read personnel data from Excel back to JSON."""

import json
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import jsonschema
except ImportError:
    jsonschema = None


def load_json(path):
    with open(path) as f:
        return json.load(f)


def write_json(path, data):
    with open(path, "w") as f:
        json.dump(data, f, indent=2)
        f.write("\n")


def collect_skill_ids(root):
    ids = set()
    taxonomy = root / "taxonomy"
    for p in taxonomy.rglob("*.json"):
        if p.name.startswith("_"):
            continue
        data = load_json(p)
        if "id" in data:
            ids.add(data["id"])
    return ids


def collect_personnel(root):
    personnel = {}
    d = root / "personnel"
    for p in sorted(d.glob("*.json")):
        if p.name.startswith("_") or p.name == "person-template.json":
            continue
        data = load_json(p)
        personnel[data["id"]] = {"path": p, "data": data}
    return personnel


def import_assessments(root, wb, skill_ids, existing_personnel):
    """Import from the Assessments sheet (matrix format)."""
    ws = wb["Assessments"]
    changes = []

    # Read skill IDs from header row
    header_skills = {}
    for col in range(3, ws.max_column + 1):
        sid = ws.cell(row=1, column=col).value
        if sid and sid in skill_ids:
            header_skills[col] = sid

    for row in range(2, ws.max_row + 1):
        person_id = ws.cell(row=row, column=1).value
        if not person_id or person_id not in existing_personnel:
            continue

        person = existing_personnel[person_id]["data"]
        skill_map = {s["skill"]: s for s in person.get("skills", [])}

        for col, sid in header_skills.items():
            cell_value = ws.cell(row=row, column=col).value
            if cell_value is None:
                continue
            try:
                level = int(cell_value)
            except (ValueError, TypeError):
                continue

            if sid in skill_map:
                if skill_map[sid]["level"] != level:
                    changes.append(
                        f"  {person_id}: {sid} level {skill_map[sid]['level']} → {level}"
                    )
                    skill_map[sid]["level"] = level
            else:
                changes.append(f"  {person_id}: {sid} NEW at level {level}")
                person.setdefault("skills", []).append(
                    {
                        "skill": sid,
                        "level": level,
                        "assessed_date": "2026-01-01",
                        "assessed_by": "IMPORT",
                    }
                )

    return changes


def import_details(root, wb, skill_ids, existing_personnel):
    """Import from the Skill Details sheet."""
    ws = wb["Skill Details"]
    changes = []

    for row in range(2, ws.max_row + 1):
        person_id = ws.cell(row=row, column=1).value
        if not person_id or person_id not in existing_personnel:
            continue

        skill_id = ws.cell(row=row, column=3).value
        if not skill_id or skill_id not in skill_ids:
            continue

        person = existing_personnel[person_id]["data"]
        skill_map = {s["skill"]: s for s in person.get("skills", [])}

        level = ws.cell(row=row, column=4).value
        assessed_date = ws.cell(row=row, column=5).value
        assessed_by = ws.cell(row=row, column=6).value
        last_practiced = ws.cell(row=row, column=7).value
        evidence = ws.cell(row=row, column=8).value
        authorized = ws.cell(row=row, column=9).value
        auth_date = ws.cell(row=row, column=10).value
        authorizer = ws.cell(row=row, column=11).value
        auth_scope = ws.cell(row=row, column=12).value

        entry = skill_map.get(skill_id)
        if entry is None:
            entry = {
                "skill": skill_id,
                "level": int(level) if level else 0,
                "assessed_date": str(assessed_date) if assessed_date else "2026-01-01",
                "assessed_by": str(assessed_by) if assessed_by else "IMPORT",
            }
            person.setdefault("skills", []).append(entry)
            changes.append(f"  {person_id}: {skill_id} NEW from details")
        else:
            if level is not None:
                entry["level"] = int(level)

        if assessed_date:
            entry["assessed_date"] = str(assessed_date)[:10]
        if assessed_by:
            entry["assessed_by"] = str(assessed_by)
        if last_practiced:
            entry["last_practiced"] = str(last_practiced)[:10]
        if evidence:
            entry["evidence_notes"] = str(evidence)

        if authorized:
            entry["authorization"] = {
                "authorized": bool(authorized),
                "date": str(auth_date)[:10] if auth_date else "",
                "authorizer": str(authorizer) if authorizer else "",
                "scope": str(auth_scope) if auth_scope else "",
            }

    return changes


def import_profiles(root, wb, existing_personnel):
    """Import profile assignments from the Profiles sheet."""
    ws = wb["Profiles"]
    changes = []

    for row in range(2, ws.max_row + 1):
        person_id = ws.cell(row=row, column=1).value
        if not person_id or person_id not in existing_personnel:
            continue

        profiles_str = ws.cell(row=row, column=3).value
        if profiles_str is None:
            continue

        new_profiles = [p.strip() for p in str(profiles_str).split(",") if p.strip()]
        person = existing_personnel[person_id]["data"]
        old_profiles = person.get("profiles", [])

        if set(new_profiles) != set(old_profiles):
            changes.append(f"  {person_id}: profiles {old_profiles} → {new_profiles}")
            person["profiles"] = new_profiles

    return changes


def import_dev_plans(root, wb, skill_ids, existing_personnel):
    """Import development plans from the Development Plans sheet."""
    ws = wb["Development Plans"]
    changes = []

    # Group by person
    plans_by_person = {}
    for row in range(2, ws.max_row + 1):
        person_id = ws.cell(row=row, column=1).value
        if not person_id or person_id not in existing_personnel:
            continue

        skill_id = ws.cell(row=row, column=3).value
        if not skill_id or skill_id not in skill_ids:
            continue

        entry = {
            "skill": skill_id,
            "current_level": int(ws.cell(row=row, column=4).value or 0),
            "target_level": int(ws.cell(row=row, column=5).value or 1),
            "target_date": str(ws.cell(row=row, column=6).value or "")[:10],
        }
        actions_str = ws.cell(row=row, column=7).value
        if actions_str:
            entry["actions"] = [a.strip() for a in str(actions_str).split(";") if a.strip()]

        plans_by_person.setdefault(person_id, []).append(entry)

    for person_id, plans in plans_by_person.items():
        person = existing_personnel[person_id]["data"]
        old_count = len(person.get("development_plan", []))
        person["development_plan"] = plans
        changes.append(
            f"  {person_id}: development_plan updated ({old_count} → {len(plans)} targets)"
        )

    return changes


def run_import(root, args):
    """Entry point from CLI."""
    if load_workbook is None:
        print("ERROR: openpyxl not installed")
        return

    xlsx_path = Path(args.file)
    if not xlsx_path.exists():
        print(f"ERROR: File not found: {xlsx_path}")
        return

    wb = load_workbook(xlsx_path)
    skill_ids = collect_skill_ids(root)
    existing = collect_personnel(root)

    if not existing:
        print("No existing personnel records found in personnel/")
        return

    all_changes = []

    if "Assessments" in wb.sheetnames:
        all_changes.extend(import_assessments(root, wb, skill_ids, existing))
    if "Skill Details" in wb.sheetnames:
        all_changes.extend(import_details(root, wb, skill_ids, existing))
    if "Profiles" in wb.sheetnames:
        all_changes.extend(import_profiles(root, wb, existing))
    if "Development Plans" in wb.sheetnames:
        all_changes.extend(import_dev_plans(root, wb, skill_ids, existing))

    if not all_changes:
        print("No changes detected.")
        return

    print(f"Changes detected ({len(all_changes)}):")
    for change in all_changes:
        print(change)

    # Validate resulting records against person schema
    schema_path = root / "schema" / "person.schema.json"
    validation_errors = []
    if schema_path.exists() and jsonschema is not None:
        schema = load_json(schema_path)
        for person_id, info in existing.items():
            try:
                jsonschema.validate(info["data"], schema)
            except jsonschema.ValidationError as e:
                validation_errors.append(f"  {person_id}: {e.message} (at {e.json_path})")
    elif jsonschema is None:
        print("\nWARNING: jsonschema not installed, skipping schema validation")

    if validation_errors:
        print(f"\nSchema validation FAILED ({len(validation_errors)} error(s)):")
        for err in validation_errors:
            print(err)
        print("\nImport aborted — fix the data and retry.")
        return

    if args.dry_run:
        print("\nDry run — no files written.")
        return

    confirm = input(f"\nWrite {len(all_changes)} changes? [y/N]: ")
    if confirm.lower() != "y":
        print("Aborted.")
        return

    for person_id, info in existing.items():
        write_json(info["path"], info["data"])
        print(f"WRITTEN: {info['path']}")

    print(f"\nImport complete. {len(all_changes)} changes applied.")
