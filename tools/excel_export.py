"""Excel export module: produce workbooks from JSON data."""

import json

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


def load_json(path):
    with open(path) as f:
        return json.load(f)


def collect_files(root, directory):
    d = root / directory
    if not d.exists():
        return []
    return sorted(
        p
        for p in d.rglob("*.json")
        if not p.name.startswith("_") and p.name != "person-template.json"
    )


HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def style_header(ws, row=1, max_col=None):
    if max_col is None:
        max_col = ws.max_column
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def auto_width(ws, min_width=10, max_width=50):
    for col in ws.columns:
        length = min_width
        for cell in col:
            if cell.value:
                length = max(length, min(len(str(cell.value)), max_width))
        ws.column_dimensions[get_column_letter(col[0].column)].width = length + 2


# ---------- Taxonomy Export ----------


def export_taxonomy(root, output_path):
    """Export skill taxonomy to Excel workbook."""
    if Workbook is None:
        print("ERROR: openpyxl not installed")
        return
    wb = Workbook()
    wb.remove(wb.active)

    meta = load_json(root / "taxonomy" / "_meta.json")
    skills_by_domain = {}

    for path in collect_files(root, "taxonomy"):
        data = load_json(path)
        domain = data.get("domain", "unknown")
        skills_by_domain.setdefault(domain, []).append(data)

    for domain_info in meta["domains"]:
        domain_dir = domain_info["directory"]
        domain_name = domain_info["name"]
        skills = skills_by_domain.get(domain_dir, [])
        if not skills:
            continue

        ws = wb.create_sheet(title=domain_name[:31])
        headers = [
            "ID",
            "Name",
            "Category",
            "Status",
            "Description",
            "Requires",
            "Recommended",
            "Equipment",
            "Tags",
        ]
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        style_header(ws, max_col=len(headers))

        for row_idx, skill in enumerate(sorted(skills, key=lambda s: s["id"]), 2):
            ws.cell(row=row_idx, column=1, value=skill["id"])
            ws.cell(row=row_idx, column=2, value=skill["name"])
            ws.cell(row=row_idx, column=3, value=skill.get("category", ""))
            ws.cell(row=row_idx, column=4, value=skill.get("status", ""))
            ws.cell(
                row=row_idx, column=5, value=skill.get("description", "")
            ).alignment = WRAP_ALIGN
            ws.cell(row=row_idx, column=6, value=", ".join(skill.get("requires", [])))
            ws.cell(row=row_idx, column=7, value=", ".join(skill.get("recommended", [])))
            ws.cell(row=row_idx, column=8, value=", ".join(skill.get("equipment", [])))
            ws.cell(row=row_idx, column=9, value=", ".join(skill.get("tags", [])))

        auto_width(ws)
        ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)
    print(f"EXPORTED: {output_path} ({sum(len(v) for v in skills_by_domain.values())} skills)")


# ---------- Personnel Export ----------


def export_personnel(root, output_path):
    """Export personnel records to Excel workbook (round-trip format)."""
    if Workbook is None:
        print("ERROR: openpyxl not installed")
        return
    wb = Workbook()

    # Collect all skill IDs for column headers
    all_skill_ids = []
    for path in collect_files(root, "taxonomy"):
        data = load_json(path)
        all_skill_ids.append(data["id"])
    all_skill_ids.sort()

    # Collect personnel
    personnel = []
    for path in collect_files(root, "personnel"):
        personnel.append(load_json(path))
    personnel.sort(key=lambda p: p["id"])

    # --- Assessments sheet ---
    ws_assess = wb.active
    ws_assess.title = "Assessments"

    # Header row 1: skill IDs
    ws_assess.cell(row=1, column=1, value="Person ID")
    ws_assess.cell(row=1, column=2, value="Name")
    col = 3
    skill_col_map = {}
    for sid in all_skill_ids:
        ws_assess.cell(row=1, column=col, value=sid)
        skill_col_map[sid] = col
        col += 1
    style_header(ws_assess, max_col=col - 1)

    # Precompute detail row index: (person_id, skill_id) -> detail row number
    detail_row_map = {}
    detail_row = 2
    for person in personnel:
        for entry in person.get("skills", []):
            detail_row_map[(person["id"], entry["skill"])] = detail_row
            detail_row += 1

    # Data rows: person × skill levels (with hyperlinks to detail sheet)
    link_font = Font(color="0563C1", underline="single")
    for row_idx, person in enumerate(personnel, 2):
        ws_assess.cell(row=row_idx, column=1, value=person["id"])
        ws_assess.cell(row=row_idx, column=2, value=person["name"])
        skill_map = {s["skill"]: s for s in person.get("skills", [])}
        for sid, c in skill_col_map.items():
            entry = skill_map.get(sid)
            if entry:
                cell = ws_assess.cell(row=row_idx, column=c, value=entry["level"])
                dr = detail_row_map.get((person["id"], sid))
                if dr:
                    cell.hyperlink = f"#'Skill Details'!A{dr}"
                    cell.font = link_font

    ws_assess.freeze_panes = "C2"
    auto_width(ws_assess, min_width=5, max_width=20)

    # --- Details sheet ---
    ws_detail = wb.create_sheet("Skill Details")
    detail_headers = [
        "Person ID",
        "Name",
        "Skill",
        "Level",
        "Assessed Date",
        "Assessed By",
        "Last Practiced",
        "Evidence Notes",
        "Authorized",
        "Auth Date",
        "Authorizer",
        "Auth Scope",
    ]
    for col, h in enumerate(detail_headers, 1):
        ws_detail.cell(row=1, column=col, value=h)
    style_header(ws_detail, max_col=len(detail_headers))

    row_idx = 2
    for person in personnel:
        for entry in person.get("skills", []):
            ws_detail.cell(row=row_idx, column=1, value=person["id"])
            ws_detail.cell(row=row_idx, column=2, value=person["name"])
            ws_detail.cell(row=row_idx, column=3, value=entry["skill"])
            ws_detail.cell(row=row_idx, column=4, value=entry["level"])
            ws_detail.cell(row=row_idx, column=5, value=entry.get("assessed_date", ""))
            ws_detail.cell(row=row_idx, column=6, value=entry.get("assessed_by", ""))
            ws_detail.cell(row=row_idx, column=7, value=entry.get("last_practiced", ""))
            ws_detail.cell(row=row_idx, column=8, value=entry.get("evidence_notes", ""))
            auth = entry.get("authorization")
            if auth and isinstance(auth, dict):
                ws_detail.cell(row=row_idx, column=9, value=auth.get("authorized", False))
                ws_detail.cell(row=row_idx, column=10, value=auth.get("date", ""))
                ws_detail.cell(row=row_idx, column=11, value=auth.get("authorizer", ""))
                ws_detail.cell(row=row_idx, column=12, value=auth.get("scope", ""))
            row_idx += 1

    auto_width(ws_detail)
    ws_detail.auto_filter.ref = ws_detail.dimensions

    # --- Development Plans sheet ---
    ws_dev = wb.create_sheet("Development Plans")
    dev_headers = [
        "Person ID",
        "Name",
        "Skill",
        "Current Level",
        "Target Level",
        "Target Date",
        "Actions",
    ]
    for col, h in enumerate(dev_headers, 1):
        ws_dev.cell(row=1, column=col, value=h)
    style_header(ws_dev, max_col=len(dev_headers))

    row_idx = 2
    for person in personnel:
        for target in person.get("development_plan", []):
            ws_dev.cell(row=row_idx, column=1, value=person["id"])
            ws_dev.cell(row=row_idx, column=2, value=person["name"])
            ws_dev.cell(row=row_idx, column=3, value=target["skill"])
            ws_dev.cell(row=row_idx, column=4, value=target["current_level"])
            ws_dev.cell(row=row_idx, column=5, value=target["target_level"])
            ws_dev.cell(row=row_idx, column=6, value=target.get("target_date", ""))
            ws_dev.cell(row=row_idx, column=7, value="; ".join(target.get("actions", [])))
            row_idx += 1

    auto_width(ws_dev)

    # --- Profiles sheet ---
    ws_prof = wb.create_sheet("Profiles")
    prof_headers = ["Person ID", "Name", "Assigned Profiles"]
    for col, h in enumerate(prof_headers, 1):
        ws_prof.cell(row=1, column=col, value=h)
    style_header(ws_prof, max_col=len(prof_headers))

    for row_idx, person in enumerate(personnel, 2):
        ws_prof.cell(row=row_idx, column=1, value=person["id"])
        ws_prof.cell(row=row_idx, column=2, value=person["name"])
        ws_prof.cell(row=row_idx, column=3, value=", ".join(person.get("profiles", [])))

    auto_width(ws_prof)

    wb.save(output_path)
    print(f"EXPORTED: {output_path} ({len(personnel)} personnel, {len(all_skill_ids)} skills)")


# ---------- Reports Export ----------


def export_reports(root, output_path):
    """Export gap analysis and service readiness reports to Excel."""
    if Workbook is None:
        print("ERROR: openpyxl not installed")
        return

    from gap_analysis import compute_gap_analysis, compute_service_readiness

    wb = Workbook()

    # --- Gap Analysis sheet ---
    ws_gap = wb.active
    ws_gap.title = "Gap Analysis"
    gap_headers = [
        "Person",
        "Profile",
        "Skill",
        "Skill Name",
        "Required Level",
        "Current Level",
        "Gap",
        "Criticality",
    ]
    for col, h in enumerate(gap_headers, 1):
        ws_gap.cell(row=1, column=col, value=h)
    style_header(ws_gap, max_col=len(gap_headers))

    gap_red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    gap_yellow = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    gaps = compute_gap_analysis(root)
    row_idx = 2
    for entry in gaps:
        ws_gap.cell(row=row_idx, column=1, value=entry["person"])
        ws_gap.cell(row=row_idx, column=2, value=entry["profile"])
        ws_gap.cell(row=row_idx, column=3, value=entry["skill"])
        ws_gap.cell(row=row_idx, column=4, value=entry.get("skill_name", ""))
        ws_gap.cell(row=row_idx, column=5, value=entry["required_level"])
        ws_gap.cell(row=row_idx, column=6, value=entry["current_level"])
        ws_gap.cell(row=row_idx, column=7, value=entry["gap"])
        ws_gap.cell(row=row_idx, column=8, value=entry["criticality"])
        if entry["gap"] > 0:
            fill = gap_red if entry["criticality"] == "essential" else gap_yellow
            for c in range(1, len(gap_headers) + 1):
                ws_gap.cell(row=row_idx, column=c).fill = fill
        row_idx += 1

    auto_width(ws_gap)
    ws_gap.auto_filter.ref = ws_gap.dimensions

    # --- Service Readiness sheet ---
    ws_svc = wb.create_sheet("Service Readiness")
    svc_headers = ["Service", "Skill", "Required Level", "Qualified Personnel"]
    for col, h in enumerate(svc_headers, 1):
        ws_svc.cell(row=1, column=col, value=h)
    style_header(ws_svc, max_col=len(svc_headers))

    readiness = compute_service_readiness(root)
    row_idx = 2
    for entry in readiness:
        ws_svc.cell(row=row_idx, column=1, value=entry["service"])
        ws_svc.cell(row=row_idx, column=2, value=entry["skill"])
        ws_svc.cell(row=row_idx, column=3, value=entry["required_level"])
        ws_svc.cell(row=row_idx, column=4, value=", ".join(entry["qualified_personnel"]))
        if not entry["qualified_personnel"]:
            for c in range(1, len(svc_headers) + 1):
                ws_svc.cell(row=row_idx, column=c).fill = gap_red
        row_idx += 1

    auto_width(ws_svc)
    ws_svc.auto_filter.ref = ws_svc.dimensions

    wb.save(output_path)
    print(f"EXPORTED: {output_path}")


def run_export(root, args):
    """Entry point from CLI."""
    if args.type == "taxonomy":
        export_taxonomy(root, root / "taxonomy-export.xlsx")
    elif args.type == "personnel":
        export_personnel(root, root / "personnel-export.xlsx")
    elif args.type == "reports":
        export_reports(root, root / "reports-export.xlsx")
    elif args.type == "all":
        export_taxonomy(root, root / "taxonomy-export.xlsx")
        export_personnel(root, root / "personnel-export.xlsx")
        export_reports(root, root / "reports-export.xlsx")
