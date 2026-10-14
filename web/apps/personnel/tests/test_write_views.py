import datetime
import io

import pytest
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook

from apps.accounts.models import User
from apps.personnel.models import Person, Project


@pytest.fixture
def writer(db):
    """A manager: edits personnel, assessments and projects, but may not delete
    people or manage accounts."""
    return User.objects.create_user(username="writer", password="strong-pass-here", role="manager")


@pytest.fixture
def logged_in(writer):
    client = Client()
    client.force_login(writer)
    return client


@pytest.fixture
def person(db):
    return Person.objects.create(
        id="JD01", name="Jane Doe", domains=[], start_date=datetime.date(2024, 1, 15)
    )


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestProjectStatus:
    def test_dashboard_counts_open_projects(self, logged_in):
        Project.objects.create(id="p1", name="Open One", status="planning")
        Project.objects.create(id="p2", name="Open Two", status="planning")
        Project.objects.create(id="p3", name="Finished", status="done")
        Project.objects.create(id="p4", name="Shelved", status="archived")

        response = logged_in.get("/")
        assert response.status_code == 200
        assert response.context["open_count"] == 2
        assert response.context["done_count"] == 1

    def test_project_edit_rejects_invalid_status(self, logged_in):
        Project.objects.create(id="p1", name="Open One", status="planning")
        response = logged_in.post(
            reverse("personnel:project-edit", args=["p1"]),
            {"name": "Open One", "status": "active"},
        )
        assert response.status_code == 200  # re-rendered form, not a redirect
        assert b"status must be one of" in response.content
        assert Project.objects.get(pk="p1").status == "planning"

    def test_project_edit_accepts_valid_status(self, logged_in):
        Project.objects.create(id="p1", name="Open One", status="planning")
        response = logged_in.post(
            reverse("personnel:project-edit", args=["p1"]),
            {"name": "Open One", "status": "archived"},
        )
        assert response.status_code == 302
        assert Project.objects.get(pk="p1").status == "archived"

    def test_person_detail_fitness_uses_open_projects(self, logged_in, person):
        """Completed projects are history, not staffing candidates."""
        Project.objects.create(
            id="p1",
            name="Open",
            status="planning",
            required_competencies=[{"skill": "FOUND-REPORT", "minimum_level": 3}],
        )
        Project.objects.create(
            id="p2",
            name="Finished",
            status="done",
            required_competencies=[{"skill": "FOUND-REPORT", "minimum_level": 3}],
        )
        response = logged_in.get(reverse("personnel:person-detail", args=["JD01"]))
        assert response.status_code == 200
        names = [f["project"].name for f in response.context["project_fitness"]]
        assert names == ["Open"]


@pytest.fixture
def taxonomy_domain(db):
    from apps.taxonomy.models import Domain

    return Domain.objects.using("taxonomy").create(
        id="FOUND", name="Foundations", description="Cross-cutting", categories=[]
    )


@pytest.fixture
def taxonomy_skill(taxonomy_domain):
    from apps.taxonomy.models import Skill

    return Skill.objects.using("taxonomy").create(
        id="FOUND-REPORT",
        name="Report Writing",
        domain=taxonomy_domain,
        category="foundations",
        status="active",
        description="Technical report writing",
        proficiency={},
    )


@pytest.fixture
def person_skill(person, taxonomy_skill):
    from apps.personnel.models import PersonSkill

    return PersonSkill.objects.create(
        person=person,
        skill_id="FOUND-REPORT",
        level=3,
        assessed_date=datetime.date(2024, 6, 1),
        assessed_by_id="LEAD01",
        evidence_notes="Wrote the Q2 report",
    )


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestLevelBounds:
    @pytest.mark.parametrize("bad_level", ["999", "-1", "6", "high"])
    def test_skill_edit_rejects_out_of_range(self, logged_in, person_skill, bad_level):
        response = logged_in.post(
            reverse("personnel:person-skill-edit", args=["JD01", "FOUND-REPORT"]),
            {"level": bad_level, "assessed_date": "2026-06-01"},
        )
        assert response.status_code == 200
        person_skill.refresh_from_db()
        assert person_skill.level == 3  # unchanged

    def test_skill_edit_accepts_valid_level(self, logged_in, person_skill):
        response = logged_in.post(
            reverse("personnel:person-skill-edit", args=["JD01", "FOUND-REPORT"]),
            {"level": "5", "assessed_date": "2026-06-01", "evidence_notes": "n"},
        )
        assert response.status_code == 302
        person_skill.refresh_from_db()
        assert person_skill.level == 5


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestDateValidation:
    @pytest.mark.parametrize("bad_date", ["2026-13-45", "01/06/2026", "yesterday", ""])
    def test_person_create_rejects_bad_date(self, logged_in, bad_date):
        response = logged_in.post(
            reverse("personnel:person-create"),
            {"name": "New Person", "start_date": bad_date},
        )
        assert response.status_code == 200
        assert not Person.objects.filter(name="New Person").exists()

    def test_person_create_accepts_iso_date(self, logged_in):
        response = logged_in.post(
            reverse("personnel:person-create"),
            {"name": "New Person", "start_date": "2026-06-01"},
        )
        assert response.status_code == 302
        created = Person.objects.get(name="New Person")
        assert created.start_date == datetime.date(2026, 6, 1)

    def test_person_edit_rejects_bad_date(self, logged_in, person):
        response = logged_in.post(
            reverse("personnel:person-edit", args=["JD01"]),
            {"name": "Jane Doe", "start_date": "2026-02-30"},
        )
        assert response.status_code == 200
        person.refresh_from_db()
        assert person.start_date == datetime.date(2024, 1, 15)


SHEET = "Self-Assessment"
HEADERS = ["Skill ID", "Skill Name", "Domain", "Category", "Level", "Description", "Notes"]


def _workbook(rows, sheet_title=SHEET, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    if headers is not None:
        ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "upload.xlsx"
    return buf


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestSelfAssessmentImport:
    def url(self):
        return reverse("personnel:self-assessment-import", args=["JD01"])

    def test_imports_valid_rows(self, logged_in, person, taxonomy_skill):
        from apps.personnel.models import SelfAssessment

        upload = _workbook(
            [["FOUND-REPORT", "Report Writing", "FOUND", "f", "3 — Independent", "", "did it"]]
        )
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 302
        sa = SelfAssessment.objects.get(person=person, skill_id="FOUND-REPORT")
        assert sa.level == 3
        assert sa.evidence_notes == "did it"

    def test_rejects_workbook_without_expected_sheet(self, logged_in, person, taxonomy_skill):
        upload = _workbook([["FOUND-REPORT", "x", "x", "x", 3, "", ""]], sheet_title="Budget")
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert b"Self-Assessment" in response.content
        from apps.personnel.models import SelfAssessment

        assert not SelfAssessment.objects.filter(person=person).exists()

    def test_rejects_workbook_with_wrong_headers(self, logged_in, person, taxonomy_skill):
        upload = _workbook([["FOUND-REPORT", "x", 3]], headers=["Thing", "Other", "Value"])
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert b"column" in response.content.lower()

    def test_reports_rows_it_could_not_read(self, logged_in, person, taxonomy_skill):
        """A row with an unreadable level must be named, not silently dropped."""
        upload = _workbook(
            [
                ["FOUND-REPORT", "Report Writing", "FOUND", "f", "3", "", "ok"],
                ["FOUND-REPORT-X", "Bogus", "FOUND", "f", "Level 3", "", ""],
            ]
        )
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert b"FOUND-REPORT-X" in response.content
        assert b"1 row" in response.content or b"imported 1" in response.content.lower()

    def test_rejects_file_over_the_size_limit(self, logged_in, person, settings):
        settings.MAX_IMPORT_FILE_SIZE = 128
        upload = _workbook([["FOUND-REPORT", "x", "x", "x", 3, "", ""]])
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert b"too large" in response.content.lower()

    def test_rejects_a_file_that_is_not_a_workbook(self, logged_in, person):
        upload = io.BytesIO(b"this is not a spreadsheet")
        upload.name = "notes.xlsx"
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert b"Invalid Excel file" in response.content

    def test_round_trips_the_exported_workbook(self, logged_in, person, taxonomy_skill):
        """Whatever generate_self_assessment_workbook writes must import cleanly."""
        from apps.core.export import generate_self_assessment_workbook
        from apps.personnel.models import SelfAssessment

        wb = generate_self_assessment_workbook(person)
        ws = wb["Self-Assessment"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "FOUND-REPORT":
                row[4].value = "4 — Expert"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "export.xlsx"

        response = logged_in.post(self.url(), {"file": buf})
        assert response.status_code == 302
        assert SelfAssessment.objects.get(person=person, skill_id="FOUND-REPORT").level == 4


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestProjectAssess:
    def test_preserves_existing_evidence_notes(self, logged_in, person_skill):
        Project.objects.create(
            id="p1",
            name="Audit Engagement",
            status="planning",
            required_competencies=[{"skill": "FOUND-REPORT", "minimum_level": 3}],
        )
        response = logged_in.post(
            reverse("personnel:project-assess", args=["p1"]),
            {"person_JD01_skill_FOUND-REPORT": "4"},
        )
        assert response.status_code == 302
        person_skill.refresh_from_db()
        assert person_skill.level == 4
        assert "Wrote the Q2 report" in person_skill.evidence_notes
        assert "Audit Engagement" in person_skill.evidence_notes

    def test_sets_notes_on_a_new_record(self, logged_in, person, taxonomy_skill):
        from apps.personnel.models import PersonSkill

        Project.objects.create(
            id="p1",
            name="Audit Engagement",
            status="planning",
            required_competencies=[{"skill": "FOUND-REPORT", "minimum_level": 3}],
        )
        logged_in.post(
            reverse("personnel:project-assess", args=["p1"]),
            {"person_JD01_skill_FOUND-REPORT": "2"},
        )
        ps = PersonSkill.objects.get(person=person, skill_id="FOUND-REPORT")
        assert ps.level == 2
        assert ps.evidence_notes == "Project: Audit Engagement"

    def test_does_not_record_the_same_project_twice(self, logged_in, person_skill):
        Project.objects.create(
            id="p1",
            name="Audit Engagement",
            status="planning",
            required_competencies=[{"skill": "FOUND-REPORT", "minimum_level": 3}],
        )
        url = reverse("personnel:project-assess", args=["p1"])
        logged_in.post(url, {"person_JD01_skill_FOUND-REPORT": "4"})
        logged_in.post(url, {"person_JD01_skill_FOUND-REPORT": "5"})
        person_skill.refresh_from_db()
        assert person_skill.evidence_notes.count("Audit Engagement") == 1

    def test_unchanged_level_is_not_restamped(self, logged_in, person_skill):
        """The form preselects every current level, so an untouched submit
        must not record everyone as freshly assessed through this project."""
        Project.objects.create(
            id="p1",
            name="Audit Engagement",
            status="planning",
            required_competencies=[{"skill": "FOUND-REPORT", "minimum_level": 3}],
        )
        logged_in.post(
            reverse("personnel:project-assess", args=["p1"]),
            {"person_JD01_skill_FOUND-REPORT": "3"},  # already level 3
        )
        person_skill.refresh_from_db()
        assert person_skill.assessed_date == datetime.date(2024, 6, 1)
        assert person_skill.assessed_by_id == "LEAD01"
        assert "Audit Engagement" not in person_skill.evidence_notes

    def test_skips_blank_and_dash_entries(self, logged_in, person_skill):
        Project.objects.create(
            id="p1",
            name="Audit Engagement",
            status="planning",
            required_competencies=[{"skill": "FOUND-REPORT", "minimum_level": 3}],
        )
        logged_in.post(
            reverse("personnel:project-assess", args=["p1"]),
            {"person_JD01_skill_FOUND-REPORT": "-"},
        )
        person_skill.refresh_from_db()
        assert person_skill.level == 3  # untouched


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestIdGeneration:
    def test_non_latin_name_gets_a_generated_id(self, logged_in):
        """Arabic-script names are expected input here (TIME_ZONE is Asia/Dubai);
        refusing them would force staff to misspell their own names."""
        for name in ("محمد", "Иван Петров"):
            response = logged_in.post(
                reverse("personnel:person-create"),
                {"name": name, "start_date": "2026-06-01"},
            )
            assert response.status_code == 302
        ids = sorted(
            Person.objects.filter(name__in=["محمد", "Иван Петров"]).values_list("id", flat=True)
        )
        assert ids == ["person-1", "person-2"]
        assert Person.objects.get(id="person-1").name == "محمد"

    def test_non_latin_project_name_gets_a_generated_id(self, logged_in):
        response = logged_in.post(
            reverse("personnel:project-create"), {"name": "مشروع", "status": "planning"}
        )
        assert response.status_code == 302
        assert Project.objects.get(name="مشروع").id == "project-1"

    def test_never_creates_an_empty_primary_key(self, logged_in):
        logged_in.post(
            reverse("personnel:person-create"),
            {"name": "日本語", "start_date": "2026-06-01"},
        )
        assert not Person.objects.filter(id="").exists()

    def test_long_name_produces_an_id_within_the_column(self, logged_in):
        response = logged_in.post(
            reverse("personnel:person-create"),
            {
                "name": "Jean-Philippe Alexander Vandenberghe-Fitzwilliam",
                "start_date": "2026-06-01",
            },
        )
        assert response.status_code == 302
        created = Person.objects.get(name__startswith="Jean-Philippe")
        assert 0 < len(created.id) <= 50

    def test_duplicate_names_get_distinct_ids(self, logged_in):
        for _ in range(2):
            logged_in.post(
                reverse("personnel:person-create"),
                {"name": "Sam Smith", "start_date": "2026-06-01"},
            )
        ids = sorted(Person.objects.filter(name="Sam Smith").values_list("id", flat=True))
        assert ids == ["sam-smith", "sam-smith-2"]


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestSelfAssessmentAccess:
    @pytest.fixture
    def other_person(self, db):
        return Person.objects.create(
            id="AB02",
            name="Alex Brown",
            domains=[],
            start_date=datetime.date(2024, 3, 1),
        )

    @pytest.fixture
    def linked_user(self, db, person):
        user = User.objects.create_user(username="jane", password="strong-pass-here", role="user")
        user.person = person
        user.save()
        return user

    def test_owner_may_import_their_own(self, linked_user, person, taxonomy_skill):
        from apps.personnel.models import SelfAssessment

        client = Client()
        client.force_login(linked_user)
        upload = _workbook([["FOUND-REPORT", "Report Writing", "FOUND", "f", "3", "", "mine"]])
        response = client.post(
            reverse("personnel:self-assessment-import", args=["JD01"]), {"file": upload}
        )
        assert response.status_code == 302
        assert SelfAssessment.objects.filter(person=person).exists()

    def test_owner_may_not_import_for_someone_else(self, linked_user, other_person, taxonomy_skill):
        from apps.personnel.models import SelfAssessment

        client = Client()
        client.force_login(linked_user)
        upload = _workbook([["FOUND-REPORT", "Report Writing", "FOUND", "f", "5", "", "not mine"]])
        response = client.post(
            reverse("personnel:self-assessment-import", args=["AB02"]), {"file": upload}
        )
        assert response.status_code == 403
        assert not SelfAssessment.objects.filter(person=other_person).exists()

    def test_unlinked_user_may_not_import_for_anyone(self, person, taxonomy_skill):
        client = Client()
        client.force_login(
            User.objects.create_user(username="nobody", password="strong-pass-here", role="user")
        )
        upload = _workbook([["FOUND-REPORT", "x", "x", "x", "4", "", ""]])
        response = client.post(
            reverse("personnel:self-assessment-import", args=["JD01"]), {"file": upload}
        )
        assert response.status_code == 403

    def test_manager_may_import_for_anyone(self, logged_in, other_person, taxonomy_skill):
        upload = _workbook([["FOUND-REPORT", "x", "x", "x", "4", "", ""]])
        response = logged_in.post(
            reverse("personnel:self-assessment-import", args=["AB02"]), {"file": upload}
        )
        assert response.status_code == 302

    def test_owner_may_export_their_own(self, linked_user, person, taxonomy_skill):
        client = Client()
        client.force_login(linked_user)
        response = client.get(reverse("personnel:self-assessment-export", args=["JD01"]))
        assert response.status_code == 200

    def test_owner_may_not_export_someone_else(self, linked_user, other_person):
        client = Client()
        client.force_login(linked_user)
        response = client.get(reverse("personnel:self-assessment-export", args=["AB02"]))
        assert response.status_code == 403

    def test_seeding_into_personskill_requires_a_manager(self, linked_user, person, taxonomy_skill):
        """A person may state their own view; only a manager makes it official."""
        client = Client()
        client.force_login(linked_user)
        response = client.post(reverse("personnel:seed-self-assessment", args=["JD01"]))
        assert response.status_code == 403


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestImportBounds:
    """A small upload must not expand into an unbounded read."""

    def url(self):
        return reverse("personnel:self-assessment-import", args=["JD01"])

    def test_refuses_a_workbook_that_expands_past_the_limit(
        self, logged_in, person, taxonomy_skill, settings, monkeypatch
    ):
        import openpyxl

        from apps.personnel.models import SelfAssessment

        settings.MAX_IMPORT_UNCOMPRESSED_SIZE = 1024
        opened = []
        real = openpyxl.load_workbook
        monkeypatch.setattr(
            openpyxl, "load_workbook", lambda *a, **k: opened.append(1) or real(*a, **k)
        )
        upload = _workbook([["FOUND-REPORT", "x", "x", "x", "3", "", ""]])
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert b"uncompressed" in response.content
        assert opened == []  # refused before parsing
        assert not SelfAssessment.objects.filter(person=person).exists()

    def test_refuses_more_rows_than_the_limit(self, logged_in, person, taxonomy_skill, settings):
        settings.MAX_IMPORT_ROWS = 3
        upload = _workbook([["FOUND-REPORT", "x", "x", "x", "3", "", ""] for _ in range(10)])
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert b"more than 3 rows" in response.content

    def test_truncates_a_long_error_list(self, logged_in, person, taxonomy_skill):
        upload = _workbook([[f"BOGUS-{i}", "x", "x", "x", "3", "", ""] for i in range(200)])
        response = logged_in.post(self.url(), {"file": upload})
        assert response.status_code == 200
        assert response.content.count(b"not found in taxonomy") <= 50
        assert b"more not shown" in response.content


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestSkillAdd:
    def url(self):
        return reverse("personnel:person-skill-add", args=["JD01"])

    def test_adds_a_valid_skill(self, logged_in, person, taxonomy_skill):
        from apps.personnel.models import PersonSkill

        response = logged_in.post(
            self.url(),
            {
                "skill_id": "FOUND-REPORT",
                "level": "2",
                "assessed_date": "2026-06-01",
            },
        )
        assert response.status_code == 302
        assert PersonSkill.objects.get(person=person, skill_id="FOUND-REPORT").level == 2

    @pytest.mark.parametrize(
        "data,message",
        [
            ({"skill_id": "NOPE", "level": "2", "assessed_date": "2026-06-01"}, b"does not exist"),
            (
                {"skill_id": "FOUND-REPORT", "level": "9", "assessed_date": "2026-06-01"},
                b"between 0 and 5",
            ),
            ({"skill_id": "FOUND-REPORT", "level": "2", "assessed_date": "soon"}, b"YYYY-MM-DD"),
            ({"level": "2", "assessed_date": "2026-06-01"}, b"required"),
        ],
    )
    def test_rejects_bad_input(self, logged_in, person, taxonomy_skill, data, message):
        from apps.personnel.models import PersonSkill

        response = logged_in.post(self.url(), data)
        assert response.status_code == 200
        assert message in response.content
        assert not PersonSkill.objects.filter(person=person).exists()

    def test_rejects_a_skill_already_assigned(self, logged_in, person_skill):
        from apps.personnel.models import PersonSkill

        response = logged_in.post(
            self.url(),
            {
                "skill_id": "FOUND-REPORT",
                "level": "4",
                "assessed_date": "2026-06-01",
            },
        )
        assert response.status_code == 200
        assert b"already" in response.content
        assert PersonSkill.objects.get(pk=person_skill.pk).level == 3


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestSearchableSkillPicker:
    """Long skill lists render as a searchable combobox layered over the
    native <select>, which stays in the markup as the submitted field."""

    def test_person_skill_add_modal(self, logged_in, person, taxonomy_skill):
        response = logged_in.get(
            reverse("personnel:person-skill-add", args=["JD01"]), HTTP_HX_REQUEST="true"
        )
        assert b'x-data="combobox"' in response.content
        assert b'name="skill_id"' in response.content
        assert b'<optgroup label="Foundations">' in response.content

    def test_base_layout_loads_combobox_script(self, logged_in, person, taxonomy_skill):
        response = logged_in.get(reverse("personnel:person-skill-add", args=["JD01"]))
        assert b"combobox.js" in response.content


@pytest.fixture
def skill_catalogue(taxonomy_domain):
    """Three active skills, one deprecated, across two domains."""
    from apps.taxonomy.models import Domain, Skill

    net = Domain.objects.using("taxonomy").create(
        id="NET", name="Network Security", description="", categories=[]
    )
    for sid, name, domain, status in [
        ("FOUND-REPORT", "Report Writing", taxonomy_domain, "active"),
        ("FOUND-OLD", "Legacy Reporting", taxonomy_domain, "deprecated"),
        ("NET-FW", "Firewall Review", net, "active"),
    ]:
        Skill.objects.using("taxonomy").create(
            id=sid,
            name=name,
            domain=domain,
            category="c",
            status=status,
            description="",
            proficiency={},
        )
    Skill.objects.using("taxonomy").create(
        id="NET-SCAN",
        name="Network Scanning",
        domain=net,
        category="Reconnaissance",
        status="active",
        description="",
        proficiency={},
        tags=["port-scanning", "nmap"],
    )


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestProjectSkills:
    """The project page lists every skill; saving replaces the required set."""

    def url(self):
        return reverse("personnel:project-skills", args=["p1"])

    def make_project(self, comps):
        return Project.objects.create(
            id="p1", name="P", status="planning", required_competencies=comps
        )

    def test_save_adds_removes_and_changes_levels(self, logged_in, skill_catalogue):
        self.make_project(
            [
                {"skill": "FOUND-REPORT", "minimum_level": 3},
                {"skill": "NET-SCAN", "minimum_level": 2},
            ]
        )
        response = logged_in.post(
            self.url(),
            {
                "skill": ["NET-FW", "FOUND-REPORT"],
                "level_FOUND-REPORT": "4",
                "level_NET-FW": "1",
                "level_NET-SCAN": "5",
            },
        )
        assert response.status_code == 302
        # Kept requirements keep their position; new ones follow.
        assert Project.objects.get(pk="p1").required_competencies == [
            {"skill": "FOUND-REPORT", "minimum_level": 4},
            {"skill": "NET-FW", "minimum_level": 1},
        ]

    def test_save_with_nothing_ticked_clears_the_set(self, logged_in, skill_catalogue):
        self.make_project([{"skill": "NET-SCAN", "minimum_level": 2}])
        logged_in.post(self.url(), {})
        assert Project.objects.get(pk="p1").required_competencies == []

    def test_duplicate_ticks_are_collapsed(self, logged_in, skill_catalogue):
        self.make_project([])
        logged_in.post(self.url(), {"skill": ["NET-FW", "NET-FW"], "level_NET-FW": "2"})
        assert Project.objects.get(pk="p1").required_competencies == [
            {"skill": "NET-FW", "minimum_level": 2},
        ]

    @pytest.mark.parametrize(
        "level,message",
        [
            ("42", b"between 0 and 5"),
            ("three", b"must be a number"),
        ],
    )
    def test_rejects_bad_level_and_saves_nothing(self, logged_in, skill_catalogue, level, message):
        self.make_project([{"skill": "NET-SCAN", "minimum_level": 2}])
        response = logged_in.post(
            self.url(),
            {
                "skill": ["NET-SCAN", "NET-FW"],
                "level_NET-SCAN": "2",
                "level_NET-FW": level,
            },
        )
        assert response.status_code == 200
        assert message in response.content
        assert Project.objects.get(pk="p1").required_competencies == [
            {"skill": "NET-SCAN", "minimum_level": 2},
        ]

    def test_error_keeps_the_submitted_ticks(self, logged_in, skill_catalogue):
        self.make_project([{"skill": "NET-SCAN", "minimum_level": 2}])
        response = logged_in.post(
            self.url(),
            {
                "skill": ["NET-FW", "NOPE"],
                "level_NET-FW": "5",
            },
        )
        assert response.status_code == 200
        assert b"does not exist" in response.content
        rows = {r["id"]: r for r in response.context["skill_rows"]}
        assert rows["NET-FW"]["required"] and rows["NET-FW"]["level"] == 5
        assert not rows["NET-SCAN"]["required"]
        # The saved set is what the change bar compares against.
        assert rows["NET-SCAN"]["saved_level"] == 2
        assert rows["NET-FW"]["saved_level"] is None

    def test_cannot_add_a_deprecated_skill(self, logged_in, skill_catalogue):
        self.make_project([])
        response = logged_in.post(self.url(), {"skill": ["FOUND-OLD"], "level_FOUND-OLD": "3"})
        assert response.status_code == 200
        assert Project.objects.get(pk="p1").required_competencies == []

    def test_existing_deprecated_and_missing_requirements_survive_a_save(
        self, logged_in, skill_catalogue
    ):
        comps = [
            {"skill": "FOUND-OLD", "minimum_level": 2},
            {"skill": "GONE-SKILL", "minimum_level": 4},
        ]
        self.make_project(comps)
        response = logged_in.post(
            self.url(),
            {
                "skill": ["FOUND-OLD", "GONE-SKILL"],
                "level_FOUND-OLD": "2",
                "level_GONE-SKILL": "4",
            },
        )
        assert response.status_code == 302
        assert Project.objects.get(pk="p1").required_competencies == comps

    def test_get_redirects_to_the_project(self, logged_in, skill_catalogue):
        self.make_project([])
        response = logged_in.get(self.url())
        assert response.status_code == 302


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestProjectSkillTable:
    def test_manager_page_lists_every_active_skill(self, logged_in, person, skill_catalogue):
        from apps.personnel.models import PersonSkill

        PersonSkill.objects.create(
            person=person,
            skill_id="NET-SCAN",
            level=4,
            assessed_date=datetime.date(2026, 1, 1),
            assessed_by_id="lead",
        )
        Project.objects.create(
            id="p1",
            name="P",
            status="planning",
            required_competencies=[
                {"skill": "NET-SCAN", "minimum_level": 2},
                {"skill": "FOUND-OLD", "minimum_level": 1},
                {"skill": "GONE-SKILL", "minimum_level": 4},
            ],
        )
        response = logged_in.get(reverse("personnel:project-detail", args=["p1"]))
        rows = {r["id"]: r for r in response.context["skill_rows"]}

        assert set(rows) == {"FOUND-REPORT", "FOUND-OLD", "NET-SCAN", "NET-FW", "GONE-SKILL"}
        assert rows["NET-SCAN"] == {
            "id": "NET-SCAN",
            "name": "Network Scanning",
            "domain": "Network Security",
            "category": "Reconnaissance",
            "tags": ["port-scanning", "nmap"],
            "status": "active",
            "required": True,
            "level": 2,
            "saved_level": 2,
        }
        assert rows["NET-FW"]["required"] is False and rows["NET-FW"]["level"] == 3
        assert rows["FOUND-OLD"]["status"] == "deprecated"
        assert rows["GONE-SKILL"]["status"] == "missing"
        assert response.context["skill_holders"] == {"NET-SCAN": [["Jane Doe", 4]]}
        assert b'id="project-skill-rows"' in response.content

    def test_unrequired_deprecated_skills_are_not_listed(self, logged_in, skill_catalogue):
        Project.objects.create(id="p1", name="P", status="planning")
        response = logged_in.get(reverse("personnel:project-detail", args=["p1"]))
        assert "FOUND-OLD" not in {r["id"] for r in response.context["skill_rows"]}

    def test_skill_name_opens_the_skill_in_a_new_tab(self, logged_in, skill_catalogue):
        """The name is a link, not a checkbox label.

        Clicking a skill's name used to toggle whether the project required it,
        because the name was a <label for> bound to the checkbox. Only the
        checkbox toggles now; the name opens the skill. It opens in a new tab
        so the table's unsaved checkbox and level edits survive the detour.
        """
        Project.objects.create(id="p1", name="P", status="planning", required_competencies=[])

        html = logged_in.get(reverse("personnel:project-detail", args=["p1"])).content.decode()

        assert 'x-text="row.name"' in html
        name_tag = html[html.rindex("<", 0, html.index('x-text="row.name"')) :]
        name_tag = name_tag[: name_tag.index(">") + 1]
        assert name_tag.startswith("<a "), f"skill name is not a link: {name_tag}"
        assert 'target="_blank"' in name_tag
        assert "skillUrl(row)" in name_tag
        assert ":for=" not in name_tag, "the name must not re-bind to the checkbox"

    def test_the_requirement_checkbox_names_itself(self, logged_in, skill_catalogue):
        """The <label> used to be the checkbox's accessible name; it is a link
        now, so the checkbox has to name itself or it reads as a bare
        'checkbox' to a screen reader."""
        Project.objects.create(id="p1", name="P", status="planning", required_competencies=[])

        html = logged_in.get(reverse("personnel:project-detail", args=["p1"])).content.decode()

        box = html[html.index('type="checkbox"') :]
        box = box[: box.index(">") + 1]
        assert "aria-label" in box, f"checkbox has no accessible name: {box}"

    def test_a_row_click_opens_the_skill(self, logged_in, skill_catalogue):
        Project.objects.create(id="p1", name="P", status="planning", required_competencies=[])

        html = logged_in.get(reverse("personnel:project-detail", args=["p1"])).content.decode()

        assert "openSkill(row, $event)" in html

    def test_user_page_carries_no_skill_table_or_holder_data(self, person, skill_catalogue):
        from apps.personnel.models import PersonSkill

        PersonSkill.objects.create(
            person=person,
            skill_id="NET-SCAN",
            level=4,
            assessed_date=datetime.date(2026, 1, 1),
            assessed_by_id="lead",
        )
        Project.objects.create(
            id="p1",
            name="P",
            status="planning",
            required_competencies=[{"skill": "NET-SCAN", "minimum_level": 2}],
        )
        client = Client()
        client.force_login(
            User.objects.create_user(username="viewer", password="strong-pass-here", role="user")
        )
        response = client.get(reverse("personnel:project-detail", args=["p1"]))
        assert response.status_code == 200
        assert response.context.get("skill_rows") is None
        assert response.context.get("skill_holders") is None
        assert b"project-skill-holders" not in response.content
        assert b"Jane Doe" not in response.content


@pytest.mark.django_db(databases=["default", "taxonomy"])
class TestSkillPickerSearchesTags:
    """Skill options carry category and tags for the picker to search, not to show."""

    def test_skill_options_carry_category_and_tags(self, logged_in, person, skill_catalogue):
        response = logged_in.get(
            reverse("personnel:person-skill-add", args=["JD01"]), HTTP_HX_REQUEST="true"
        )
        assert b'data-search="Reconnaissance|port-scanning|nmap"' in response.content
